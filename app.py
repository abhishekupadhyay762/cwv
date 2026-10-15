"""
Core Web Vitals Bulk Checker
============================
Checks real-user Core Web Vitals (Chrome UX Report field data) and Lighthouse
lab scores for many URLs at once, on mobile and desktop, through Google's
PageSpeed Insights API.

Every visitor uses their own PageSpeed Insights API key. Keys and uploaded
service-account files live only in the visitor's Streamlit session; nothing is
written to disk, logged or shared.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
import threading
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from datetime import datetime, timezone
from urllib.parse import urlparse

import pandas as pd
import requests
import streamlit as st

st.set_page_config(
    page_title="Core Web Vitals Bulk Checker",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# If an audit was still running when this rerun started (the visitor clicked
# "Stop audit" or changed a setting), tell its worker threads to stop.
_previous_stop = st.session_state.pop("_stop_event", None)
if _previous_stop is not None:
    _previous_stop.set()


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

PSI_ENDPOINT = "https://www.googleapis.com/pagespeedonline/v5/runPagespeed"
BASE_CATEGORIES = ["PERFORMANCE", "ACCESSIBILITY", "BEST_PRACTICES", "SEO"]
AGENTIC_CATEGORY = "AGENTIC_BROWSING"
MAX_URLS = 200
REQUEST_TIMEOUT = 120  # seconds — a Lighthouse run can be slow
AVG_SECONDS_PER_CHECK = 20  # used only for the time estimate
KEY_PATTERN = re.compile(r"^AIza[0-9A-Za-z_\-]{35}$")
DOMAIN_LIKE = re.compile(r"^(?:[a-z0-9-]+\.)+[a-z]{2,}(?::\d+)?(?:[/?#].*)?$", re.I)

DEVICE_OPTIONS = {
    "Mobile and desktop": ["mobile", "desktop"],
    "Mobile only": ["mobile"],
    "Desktop only": ["desktop"],
}
DEVICE_LABEL = {"mobile": "Mobile", "desktop": "Desktop"}

# Field (CrUX) metrics: column -> (possible API keys, converter)
FIELD_METRICS = {
    "Field LCP (s)": (["LARGEST_CONTENTFUL_PAINT_MS"], lambda v: round(v / 1000, 2)),
    "Field INP (ms)": (
        ["INTERACTION_TO_NEXT_PAINT", "INTERACTION_TO_NEXT_PAINT_MS", "EXPERIMENTAL_INTERACTION_TO_NEXT_PAINT"],
        lambda v: int(v),
    ),
    "Field CLS": (["CUMULATIVE_LAYOUT_SHIFT_SCORE"], lambda v: round(v / 100, 2)),
    "Field FCP (s)": (["FIRST_CONTENTFUL_PAINT_MS"], lambda v: round(v / 1000, 2)),
    "Field TTFB (s)": (["EXPERIMENTAL_TIME_TO_FIRST_BYTE", "TIME_TO_FIRST_BYTE"], lambda v: round(v / 1000, 2)),
}
CORE_FIELD_COLUMNS = ["Field LCP (s)", "Field INP (ms)", "Field CLS"]

SCORE_CATEGORIES = {
    "Performance": "performance",
    "Accessibility": "accessibility",
    "Best practices": "best-practices",
    "SEO": "seo",
}

# Lab (Lighthouse) metrics: column -> (audit id, converter)
LAB_METRICS = {
    "Lab LCP (s)": ("largest-contentful-paint", lambda v: round(v / 1000, 2)),
    "Lab FCP (s)": ("first-contentful-paint", lambda v: round(v / 1000, 2)),
    "Lab TBT (ms)": ("total-blocking-time", lambda v: int(round(v))),
    "Lab CLS": ("cumulative-layout-shift", lambda v: round(v, 3)),
    "Lab Speed Index (s)": ("speed-index", lambda v: round(v / 1000, 2)),
}

DEVICE_COLUMNS = [
    "URL",
    "CWV",
    "Field data",
    *FIELD_METRICS,
    *SCORE_CATEGORIES,
    "Agentic browsing",
    *LAB_METRICS,
    "Top fixes",
    "Redirects to",
    "Status",
]

COMPARE_FIELDS = ["CWV", "Performance", "Field LCP (s)", "Field INP (ms)", "Field CLS", "Lab LCP (s)", "Lab TBT (ms)"]

SHEET_FIELDS = [
    "Performance", "Accessibility", "Best practices", "SEO", "Agentic browsing",
    "CWV", "Field data", "Field LCP (s)", "Field INP (ms)", "Field CLS", "Field FCP (s)", "Field TTFB (s)",
    "Lab LCP (s)", "Lab TBT (ms)", "Lab CLS",
]

NUMBER_FORMATS = {
    "Field LCP (s)": "{:.1f}",
    "Field INP (ms)": "{:.0f}",
    "Field CLS": "{:.2f}",
    "Field FCP (s)": "{:.1f}",
    "Field TTFB (s)": "{:.1f}",
    "Performance": "{:.0f}",
    "Accessibility": "{:.0f}",
    "Best practices": "{:.0f}",
    "SEO": "{:.0f}",
    "Lab LCP (s)": "{:.1f}",
    "Lab FCP (s)": "{:.1f}",
    "Lab TBT (ms)": "{:.0f}",
    "Lab CLS": "{:.3f}",
    "Lab Speed Index (s)": "{:.1f}",
}

CELL_STYLES = {
    "good": "background-color: #E2F4EA; color: #0B6B3A;",
    "ni": "background-color: #FFF0D1; color: #8A5300;",
    "poor": "background-color: #FBE3E1; color: #A3241F;",
    "info": "background-color: #E6EDF8; color: #27466F;",
    "none": "color: #8B94A3;",
}
XLSX_FILLS = {"good": "E2F4EA", "ni": "FFF0D1", "poor": "FBE3E1", "info": "E6EDF8"}
XLSX_FONTS = {"good": "0B6B3A", "ni": "8A5300", "poor": "A3241F", "info": "27466F", "none": "8B94A3"}

CRUX_RATING = {"FAST": "good", "AVERAGE": "ni", "SLOW": "poor"}

LIGHTHOUSE_ERRORS = {
    "NO_FCP": "The page didn't show any content during the test, so Lighthouse couldn't measure it.",
    "FAILED_DOCUMENT_REQUEST": "Lighthouse couldn't load the page. It may be blocking Google's tester, timing out or offline.",
    "ERRORED_DOCUMENT_REQUEST": "The page returned an error status (4xx or 5xx) to Lighthouse.",
    "DNS_FAILURE": "The domain couldn't be found (DNS lookup failed).",
    "NOT_HTML": "The URL doesn't return an HTML page.",
    "PAGE_HUNG": "The page stopped responding during the test.",
    "CHROME_INTERSTITIAL_ERROR": "Chrome showed an error screen instead of the page, often an SSL or connection problem.",
    "INSECURE_DOCUMENT_REQUEST": "The page's SSL certificate isn't valid.",
}

COLUMN_HELP = {
    "CWV": "Core Web Vitals assessment from real Chrome users: passes when LCP, INP and CLS are all good at the 75th percentile.",
    "Field data": "URL = enough real-user traffic for this exact URL. Origin = not enough, so the numbers are for the whole domain. No data = no real-user data for the URL or its domain.",
    "Agentic browsing": "Lighthouse's Agentic Browsing check: how many of the applicable audits passed (accessibility tree, layout stability, llms.txt format).",
    "Top fixes": "The Lighthouse performance insights with the biggest estimated savings for this page.",
    "Redirects to": "The URL Lighthouse ended up on. Audit and link to the final URL instead of the redirecting one.",
}


# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,600;12..96,700&display=swap');
.block-container { max-width: 1180px; padding-top: 2.4rem; padding-bottom: 4rem; }
.hero { display: grid; grid-template-columns: minmax(0, 1.15fr) minmax(0, 1fr); gap: 2.75rem; align-items: end; margin: 0 0 1.75rem; }
.hero-title { font-family: 'Bricolage Grotesque', sans-serif; font-weight: 700; font-size: clamp(2.1rem, 4.4vw, 3.2rem); line-height: 1.02; letter-spacing: -0.025em; margin: 0 0 0.8rem; }
.hero-copy p { font-size: 1.05rem; line-height: 1.55; opacity: 0.78; max-width: 58ch; margin: 0; }
.thresholds { display: grid; grid-template-columns: repeat(3, minmax(0, 1fr)); gap: 1.1rem 1rem; }
.th-head { display: flex; align-items: baseline; gap: 0.45rem; margin-bottom: 0.45rem; }
.th-head b { font-family: 'Bricolage Grotesque', sans-serif; font-size: 1.2rem; font-weight: 700; }
.th-head span { font-size: 0.82rem; opacity: 0.65; }
.th-bar, .th-scale { display: grid; grid-template-columns: repeat(3, 1fr); gap: 3px; }
.th-bar i { display: block; height: 8px; border-radius: 2px; }
.th-bar .g { background: #13A35E; } .th-bar .n { background: #F0A01A; } .th-bar .p { background: #E2483D; }
.th-scale { margin-top: 0.35rem; font-size: 0.76rem; opacity: 0.75; font-variant-numeric: tabular-nums; }
.th-note { grid-column: 1 / -1; font-size: 0.78rem; opacity: 0.6; }
.pill { display: inline-flex; align-items: center; gap: 0.45rem; font-size: 0.83rem; padding: 0.2rem 0.65rem; border-radius: 999px; }
.pill::before { content: ""; width: 0.5rem; height: 0.5rem; border-radius: 50%; background: currentColor; }
.pill.ok { background: #E2F4EA; color: #0B6B3A; }
.pill.warn { background: #FFF0D1; color: #8A5300; }
.pill.idle { background: rgba(127, 127, 127, 0.13); }
.readout { display: grid; grid-template-columns: repeat(auto-fit, minmax(175px, 1fr)); border: 1px solid rgba(127, 127, 127, 0.28); border-radius: 0.6rem; overflow: hidden; margin: 0.25rem 0 1.25rem; }
.ro { padding: 1rem 1.15rem 1.05rem; box-shadow: -1px 0 0 rgba(127, 127, 127, 0.2); }
.ro-num { font-family: 'Bricolage Grotesque', sans-serif; font-size: 2.05rem; font-weight: 700; line-height: 1; font-variant-numeric: tabular-nums; }
.ro-num small { font-size: 1rem; font-weight: 600; opacity: 0.5; }
.ro-lab { font-size: 0.86rem; opacity: 0.75; margin-top: 0.4rem; }
.ro-bar { display: flex; height: 6px; border-radius: 3px; overflow: hidden; margin-top: 0.65rem; background: rgba(127, 127, 127, 0.16); }
.ro-bar span { display: block; height: 100%; }
.ro-sub { font-size: 0.76rem; opacity: 0.62; margin-top: 0.4rem; }
.footnote { font-size: 0.8rem; opacity: 0.62; line-height: 1.55; margin-top: 2.5rem; max-width: 90ch; }
@media (max-width: 860px) {
  .hero { grid-template-columns: 1fr; gap: 1.5rem; }
  .block-container { padding-top: 1.5rem; }
}
@media (max-width: 560px) {
  .thresholds { grid-template-columns: 1fr; gap: 0.9rem; }
  .th-head { margin-bottom: 0.3rem; }
}
</style>
""",
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------------------
# Helpers: URLs, keys, ratings
# ---------------------------------------------------------------------------

def redact(text: str, key: str) -> str:
    """Never let the visitor's API key appear in an error shown on screen."""
    text = re.sub(r"key=[^&\s)'\"]+", "key=[hidden]", text or "")
    return text.replace(key, "[hidden]") if key else text


def normalise_url(token: str) -> str | None:
    t = (token or "").strip().strip("\"'<>").rstrip(",")
    if not t:
        return None
    if not re.match(r"^https?://", t, re.I):
        if not DOMAIN_LIKE.match(t):
            return None
        t = "https://" + t
    host = urlparse(t).hostname or ""
    if "." not in host:
        return None
    return t


def clean_urls(tokens: list[str]) -> tuple[list[str], list[str], int]:
    urls, invalid, seen, dupes = [], [], set(), 0
    for token in tokens:
        if not token.strip():
            continue
        url = normalise_url(token)
        if url is None:
            invalid.append(token.strip())
        elif url in seen:
            dupes += 1
        else:
            seen.add(url)
            urls.append(url)
    return urls, invalid, dupes


def tokens_from_text(text: str) -> list[str]:
    return [t for t in re.split(r"[\s;]+", text or "") if t]


def tokens_from_upload(uploaded) -> list[str]:
    raw = uploaded.getvalue().decode("utf-8-sig", errors="ignore")
    if uploaded.name.lower().endswith(".csv"):
        tokens = []
        for row in csv.reader(io.StringIO(raw)):
            tokens.extend(cell.strip() for cell in row if cell.strip())
        return tokens
    return tokens_from_text(raw)


def score_rating(score) -> str:
    if score is None:
        return "none"
    return "good" if score >= 90 else "ni" if score >= 50 else "poor"


def audit_rating(score) -> str:
    if not isinstance(score, (int, float)):
        return "none"
    return "good" if score >= 0.9 else "ni" if score >= 0.5 else "poor"


def _same_url(a: str, b: str) -> bool:
    def norm(u):
        u = (u or "").split("#")[0]
        return u[:-1] if u.endswith("/") else u
    return norm(a) == norm(b)


def _is_blank(value) -> bool:
    if value is None or value == "":
        return True
    return isinstance(value, float) and math.isnan(value)


def _sentence(parts: list[str]) -> str:
    text = ", ".join(parts)
    return text[:1].upper() + text[1:] + "."


def fmt_duration(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 60}:{seconds % 60:02d}"


def estimate_minutes(checks: int, workers: int) -> int:
    return max(1, math.ceil(checks * AVG_SECONDS_PER_CHECK / max(1, workers) / 60))


# ---------------------------------------------------------------------------
# PageSpeed Insights: parsing
# ---------------------------------------------------------------------------

def parse_psi(payload: dict) -> tuple[dict, dict]:
    values: dict = {}
    ratings: dict = {}

    # 1. Field data (Chrome UX Report, 28 days, 75th percentile)
    exp = payload.get("loadingExperience") or {}
    metrics = exp.get("metrics") or {}
    source = "Origin" if metrics and exp.get("origin_fallback") else "URL"
    if not metrics:
        origin_exp = payload.get("originLoadingExperience") or {}
        if origin_exp.get("metrics"):
            metrics, source, exp = origin_exp["metrics"], "Origin", origin_exp
        else:
            source = "No data"
    values["Field data"] = source
    ratings["Field data"] = {"URL": "", "Origin": "info", "No data": "none"}[source]

    core_categories = []
    for col, (keys, convert) in FIELD_METRICS.items():
        metric = next((metrics[k] for k in keys if k in metrics), None)
        if metric and isinstance(metric.get("percentile"), (int, float)):
            values[col] = convert(metric["percentile"])
            ratings[col] = CRUX_RATING.get(metric.get("category"), "none")
            if col in CORE_FIELD_COLUMNS and metric.get("category"):
                core_categories.append(metric["category"])
        else:
            values[col] = None
            ratings[col] = "none"

    if source == "No data":
        cwv = "No data"
    elif core_categories:
        cwv = "Passed" if all(c == "FAST" for c in core_categories) else "Failed"
    elif exp.get("overall_category") in CRUX_RATING:
        cwv = "Passed" if exp["overall_category"] == "FAST" else "Failed"
    else:
        cwv = "No data"
    values["CWV"] = cwv
    ratings["CWV"] = {"Passed": "good", "Failed": "poor"}.get(cwv, "none")

    # 2. Lighthouse category scores
    lh = payload.get("lighthouseResult") or {}
    categories = lh.get("categories") or {}
    audits = lh.get("audits") or {}
    for col, cat_id in SCORE_CATEGORIES.items():
        score = (categories.get(cat_id) or {}).get("score")
        values[col] = int(round(score * 100)) if isinstance(score, (int, float)) else None
        ratings[col] = score_rating(values[col])

    # 3. Agentic browsing: reported by Lighthouse as "passed of applicable"
    agentic = categories.get("agentic-browsing")
    values["Agentic browsing"], ratings["Agentic browsing"] = None, "none"
    if agentic:
        passed = applicable = 0
        for ref in agentic.get("auditRefs") or []:
            if ref.get("group") == "hidden":
                continue
            result = audits.get(ref.get("id"), {})
            if result.get("scoreDisplayMode") in ("manual", "notApplicable", "informative"):
                continue
            applicable += 1
            score = result.get("score")
            if isinstance(score, (int, float)) and score >= 0.9:
                passed += 1
        if applicable:
            values["Agentic browsing"] = f"{passed} of {applicable}"
            ratings["Agentic browsing"] = "good" if passed == applicable else ("poor" if passed == 0 else "ni")

    # 4. Lab metrics (one Lighthouse run)
    for col, (audit_id, convert) in LAB_METRICS.items():
        result = audits.get(audit_id) or {}
        numeric = result.get("numericValue")
        values[col] = convert(numeric) if isinstance(numeric, (int, float)) else None
        ratings[col] = audit_rating(result.get("score")) if values[col] is not None else "none"

    # 5. Top performance fixes, ranked by estimated savings
    candidates = []
    for ref in (categories.get("performance") or {}).get("auditRefs") or []:
        if ref.get("group") in ("metrics", "hidden"):
            continue
        result = audits.get(ref.get("id"), {})
        score = result.get("score")
        if result.get("scoreDisplayMode") in ("notApplicable", "manual", "informative", "error"):
            continue
        if not isinstance(score, (int, float)) or score >= 0.9:
            continue
        savings = result.get("metricSavings") or {}
        weight = sum(float(savings.get(m) or 0) for m in ("LCP", "FCP", "TBT", "INP")) + float(savings.get("CLS") or 0) * 1000
        if not weight:
            weight = float((result.get("details") or {}).get("overallSavingsMs") or 0)
        title = re.sub(r"[`*]", "", result.get("title") or ref.get("id", ""))
        candidates.append((weight, score, title))
    candidates.sort(key=lambda c: (-c[0], c[1]))
    values["Top fixes"] = "; ".join(t for _, _, t in candidates[:3])
    ratings["Top fixes"] = ""

    # 6. Redirects
    requested = lh.get("requestedUrl") or ""
    final = lh.get("finalDisplayedUrl") or lh.get("finalUrl") or ""
    values["Redirects to"] = final if final and requested and not _same_url(final, requested) else ""
    ratings["Redirects to"] = "info" if values["Redirects to"] else ""

    return values, ratings


# ---------------------------------------------------------------------------
# PageSpeed Insights: requests and errors
# ---------------------------------------------------------------------------

def classify_error(status: int, payload) -> tuple[str, str]:
    """Turn an API error into (kind, plain-English message).

    kind: fatal (stop the whole run), rate (wait and retry), retry (temporary),
    page (problem with this URL only), category (agentic category unsupported).
    """
    err = payload.get("error") if isinstance(payload, dict) else None
    err = err if isinstance(err, dict) else {}
    message = str(err.get("message") or "")
    reasons = []
    for item in list(err.get("details") or []) + list(err.get("errors") or []):
        if isinstance(item, dict):
            reasons.append(str(item.get("reason") or ""))
    text = f"{message} {' '.join(reasons)}".lower()

    if "api key not valid" in text or "api_key_invalid" in text:
        return "fatal", "Google rejected this API key. Check you copied the whole key: it starts with AIza and is 39 characters long."
    if "referer" in text or "referrer" in text:
        return "fatal", ("This key only works on certain websites (HTTP referrer restriction), so this app's requests are blocked. "
                         "In Google Cloud Console > Credentials, set Application restrictions to None and keep the API restriction on PageSpeed Insights API.")
    if "ip_address_blocked" in text or ("ip address" in text and "blocked" in text):
        return "fatal", "This key is restricted to specific IP addresses, so this app's requests are blocked. Set Application restrictions to None for this key."
    if "service_disabled" in text or "has not been used in project" in text or "it is disabled" in text:
        return "fatal", ("The PageSpeed Insights API isn't turned on for this key's Google Cloud project. "
                         "Enable it in Google Cloud Console > APIs & Services > Library > PageSpeed Insights API, wait a couple of minutes, then run again.")
    if "api_key_service_blocked" in text or "are blocked" in text:
        return "fatal", "This key isn't allowed to call the PageSpeed Insights API. In Google Cloud Console > Credentials, add PageSpeed Insights API to the key's API restrictions."
    if status == 429 or "quota" in text or "rate_limit" in text or "resource_exhausted" in text:
        if "per day" in text or "perday" in text:
            return "fatal", "This key has used up today's PageSpeed Insights quota. Finished results are kept below; check the rest tomorrow or use another key."
        return "rate", "Google is rate-limiting this key."
    if "category" in text and ("invalid" in text or "agentic" in text):
        return "category", message
    code = re.search(r"lighthouse returned error:\s*([A-Z_]+)", message, re.I)
    if code:
        name = code.group(1).upper()
        return "page", LIGHTHOUSE_ERRORS.get(name, f"Lighthouse couldn't test this page ({name}).")
    if status >= 500:
        return "retry", "PageSpeed Insights had a temporary server error."
    if status == 400 and "url" in text:
        return "page", "PageSpeed Insights couldn't use this URL. Check it's public and typed correctly."
    return "page", message or f"PageSpeed Insights returned an error (HTTP {status})."


def run_check(url: str, strategy: str, key: str, stop_event: threading.Event, shared: dict) -> dict:
    """One PageSpeed Insights request (with retries). Runs in a worker thread."""
    base = {"url": url, "strategy": strategy}
    limits = {"retry": 2, "rate": 3, "page": 1}
    delays = {"retry": 5, "rate": 20, "page": 3}
    attempts = {k: 0 for k in limits}
    while True:
        if stop_event.is_set():
            return {**base, "error": "Not checked (audit stopped).", "skipped": True}
        categories = BASE_CATEGORIES + ([AGENTIC_CATEGORY] if shared.get("agentic", True) else [])
        params = [("url", url), ("key", key), ("strategy", strategy.upper())] + [("category", c) for c in categories]
        try:
            response = requests.get(PSI_ENDPOINT, params=params, timeout=REQUEST_TIMEOUT)
        except requests.Timeout:
            kind, message = "retry", "PageSpeed Insights didn't respond in time."
        except requests.RequestException as exc:
            kind, message = "retry", f"Couldn't reach PageSpeed Insights: {exc}"
        else:
            try:
                payload = response.json()
            except ValueError:
                payload = {}
            if response.ok and isinstance(payload, dict) and payload.get("lighthouseResult"):
                values, ratings = parse_psi(payload)
                return {**base, "values": values, "ratings": ratings}
            kind, message = classify_error(response.status_code, payload)
        message = redact(message, key)

        if kind == "category" and AGENTIC_CATEGORY in categories:
            shared["agentic"] = False  # older API without the agentic category
            continue
        if kind == "fatal":
            return {**base, "error": message, "fatal": message}
        if kind in limits and attempts[kind] < limits[kind]:
            attempts[kind] += 1
            stop_event.wait(delays[kind] * attempts[kind])
            continue
        if kind == "rate":
            message = "Google kept rate-limiting this key. Run again with fewer checks at a time."
        return {**base, "error": message}


def new_audit(urls: list[str], strategies: list[str]) -> dict:
    return {
        "urls": urls,
        "strategies": strategies,
        "results": {},
        "started": datetime.now(timezone.utc),
        "duration": None,
        "finished": False,
        "stopped": False,
        "fatal": None,
    }


def execute_audit(audit: dict, key: str, workers: int, prefix: str) -> None:
    """Run every (URL, device) check in parallel and store results in `audit`."""
    jobs = [(u, s) for u in audit["urls"] for s in audit["strategies"]]
    total = len(jobs)
    stop_event = threading.Event()
    st.session_state["_stop_event"] = stop_event
    shared = {"agentic": True}
    start = time.time()
    done = 0

    placeholder = st.empty()
    with placeholder.container(border=True):
        progress = st.progress(0.0, text=f"Starting {total} checks…")
        info_col, stop_col = st.columns([4, 1], vertical_alignment="center")
        info_col.caption(
            f"Running {min(workers, total)} checks at a time; each takes about 10 to 40 seconds. "
            "Changing anything on the page stops the audit, and finished results are kept."
        )
        stop_col.button("Stop audit", key=f"{prefix}_stop", width="stretch")

    executor = ThreadPoolExecutor(max_workers=workers, thread_name_prefix="psi")
    job_of = {executor.submit(run_check, u, s, key, stop_event, shared): (u, s) for u, s in jobs}
    pending = set(job_of)
    try:
        while pending:
            finished, pending = wait(pending, timeout=1.0, return_when=FIRST_COMPLETED)
            for future in finished:
                url, strategy = job_of[future]
                try:
                    result = future.result()
                except Exception as exc:  # parsing bug or unexpected response
                    result = {"url": url, "strategy": strategy, "error": redact(f"Unexpected error: {exc}", key)}
                audit["results"][(url, strategy)] = result
                done += 1
                if result.get("fatal") and not audit["fatal"]:
                    audit["fatal"] = result["fatal"]
                    stop_event.set()
            progress.progress(done / total, text=f"Checked {done} of {total}  ({fmt_duration(time.time() - start)} elapsed)")
            if stop_event.is_set():
                break
        audit["finished"] = audit["fatal"] is None and done == total
    finally:
        if done < total and not audit["fatal"]:
            audit["stopped"] = True
        audit["duration"] = time.time() - start
        stop_event.set()
        executor.shutdown(wait=False, cancel_futures=True)
        st.session_state.pop("_stop_event", None)
    placeholder.empty()


# ---------------------------------------------------------------------------
# Results: tables
# ---------------------------------------------------------------------------

def device_frames(audit: dict, strategy: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows, rates = [], []
    for url in audit["urls"]:
        result = audit["results"].get((url, strategy))
        row = {c: None for c in DEVICE_COLUMNS}
        rate = {c: "" for c in DEVICE_COLUMNS}
        row["URL"] = url
        if result is None:
            row["Status"], rate["Status"] = "Not checked", "none"
        elif result.get("error"):
            row["Status"], rate["Status"] = result["error"], "poor"
        else:
            row.update(result["values"])
            rate.update(result["ratings"])
            row["Status"] = "OK"
        rows.append(row)
        rates.append(rate)
    values = pd.DataFrame(rows, columns=DEVICE_COLUMNS)
    ratings = pd.DataFrame(rates, columns=DEVICE_COLUMNS)
    for col in ("CWV", "Field data", "Agentic browsing"):
        values[col] = values[col].map(lambda v: "—" if _is_blank(v) else v)
    for col in ("Top fixes", "Redirects to"):
        values[col] = values[col].map(lambda v: "" if _is_blank(v) else v)
    # Hide columns that are empty for every URL
    for col in ("Agentic browsing", "Redirects to", "Top fixes"):
        if values[col].map(lambda v: _is_blank(v) or v == "—").all():
            values, ratings = values.drop(columns=col), ratings.drop(columns=col)
    return values, ratings


def compare_frames(audit: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    per_device = {s: device_frames(audit, s) for s in audit["strategies"]}
    values = pd.DataFrame({"URL": audit["urls"]})
    ratings = pd.DataFrame({"URL": [""] * len(audit["urls"])})
    for field in COMPARE_FIELDS:
        for s in audit["strategies"]:
            v, r = per_device[s]
            if field in v:
                col = f"{field} · {DEVICE_LABEL[s]}"
                values[col] = v[field].values
                ratings[col] = r[field].values
    return values, ratings


def styled(values: pd.DataFrame, ratings: pd.DataFrame):
    formats = {}
    for col in values.columns:
        base = col.split(" · ")[0]
        if base in NUMBER_FORMATS:
            formats[col] = NUMBER_FORMATS[base]
    return (
        values.style
        .apply(lambda _: ratings[values.columns].map(lambda r: CELL_STYLES.get(r, "")), axis=None)
        .format(formats, na_rep="—")
    )


def column_config(columns) -> dict:
    config = {
        "URL": st.column_config.LinkColumn("URL", width="large"),
        "Redirects to": st.column_config.LinkColumn("Redirects to", width="medium", help=COLUMN_HELP["Redirects to"]),
        "Top fixes": st.column_config.TextColumn("Top fixes", width="large", help=COLUMN_HELP["Top fixes"]),
        "Status": st.column_config.TextColumn("Status", width="medium"),
    }
    for col in columns:
        base = col.split(" · ")[0]
        if col not in config and base in COLUMN_HELP:
            config[col] = st.column_config.Column(col, help=COLUMN_HELP[base])
    return {k: v for k, v in config.items() if k in columns}


def show_table(values: pd.DataFrame, ratings: pd.DataFrame) -> None:
    height = min(35 * (len(values) + 1) + 3, 640)
    st.dataframe(
        styled(values, ratings),
        column_config=column_config(list(values.columns)),
        hide_index=True,
        width="stretch",
        height=height,
    )


def collect_flags(audit: dict) -> dict:
    flags = {"errors": [], "origin": [], "none": [], "redirects": {}}
    for (url, strategy), result in audit["results"].items():
        device = DEVICE_LABEL[strategy]
        if result.get("error"):
            if not result.get("skipped"):
                flags["errors"].append({"URL": url, "Device": device, "Problem": result["error"]})
            continue
        source = result["values"].get("Field data")
        if source == "Origin":
            flags["origin"].append({"URL": url, "Device": device})
        elif source == "No data":
            flags["none"].append({"URL": url, "Device": device})
        if result["values"].get("Redirects to"):
            flags["redirects"][url] = result["values"]["Redirects to"]
    return flags


def summary_html(audit: dict) -> str:
    n = len(audit["urls"])
    blocks = [f"<div class='ro'><div class='ro-num'>{n}</div><div class='ro-lab'>URLs in this audit</div></div>"]
    for strategy in audit["strategies"]:
        counts = {"Passed": 0, "Failed": 0, "No data": 0, "Error": 0, "Not checked": 0}
        perf = []
        for url in audit["urls"]:
            result = audit["results"].get((url, strategy))
            if result is None:
                counts["Not checked"] += 1
            elif result.get("error"):
                counts["Error"] += 1
            else:
                counts[result["values"]["CWV"]] += 1
                if result["values"].get("Performance") is not None:
                    perf.append(result["values"]["Performance"])
        device = DEVICE_LABEL[strategy]
        segments = "".join(
            f"<span style='width:{counts[k] / n * 100:.2f}%;background:{color}'></span>"
            for k, color in (("Passed", "#13A35E"), ("Failed", "#E2483D"), ("No data", "#A7AFBC"))
            if counts[k]
        )
        detail = [f"{counts['Failed']} failed", f"{counts['No data']} without real-user data"]
        if counts["Error"]:
            detail.append(f"{counts['Error']} error{'s' if counts['Error'] > 1 else ''}")
        if counts["Not checked"]:
            detail.append(f"{counts['Not checked']} not checked")
        blocks.append(
            f"<div class='ro'><div class='ro-num'>{counts['Passed']}<small> / {n}</small></div>"
            f"<div class='ro-lab'>{device}: passed Core Web Vitals</div>"
            f"<div class='ro-bar'>{segments}</div><div class='ro-sub'>{', '.join(detail)}</div></div>"
        )
        avg = round(sum(perf) / len(perf)) if perf else None
        rating = score_rating(avg)
        color = {"good": "#0B6B3A", "ni": "#8A5300", "poor": "#A3241F"}.get(rating, "inherit")
        blocks.append(
            f"<div class='ro'><div class='ro-num' style='color:{color}'>{avg if avg is not None else '—'}</div>"
            f"<div class='ro-lab'>{device}: average performance score</div>"
            f"<div class='ro-sub'>Lighthouse lab score, 0 to 100</div></div>"
        )
    return "<div class='readout'>" + "".join(blocks) + "</div>"


# ---------------------------------------------------------------------------
# Results: downloads
# ---------------------------------------------------------------------------

def _cell(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value.item() if hasattr(value, "item") else value


def build_excel(audit: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="172033")
    wb = Workbook()

    def write_sheet(ws, values: pd.DataFrame, ratings: pd.DataFrame | None, wide: tuple = ()):
        ws.append(list(values.columns))
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for i in range(len(values)):
            ws.append([_cell(v) for v in values.iloc[i].tolist()])
            if ratings is None:
                continue
            for j, col in enumerate(values.columns, start=1):
                rating = ratings.iloc[i][col]
                cell = ws.cell(row=i + 2, column=j)
                if rating in XLSX_FILLS:
                    cell.fill = PatternFill("solid", fgColor=XLSX_FILLS[rating])
                if rating in XLSX_FONTS:
                    cell.font = Font(color=XLSX_FONTS[rating])
        for j, col in enumerate(values.columns, start=1):
            ws.column_dimensions[get_column_letter(j)].width = 58 if col in wide else max(12, min(24, len(col) + 3))
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Core Web Vitals audit"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Checked (UTC)", audit["started"].strftime("%Y-%m-%d %H:%M")])
    ws.append(["URLs", len(audit["urls"])])
    ws.append(["Devices", ", ".join(DEVICE_LABEL[s] for s in audit["strategies"])])
    for strategy in audit["strategies"]:
        values, _ = device_frames(audit, strategy)
        ok = values[values["Status"] == "OK"]
        device = DEVICE_LABEL[strategy]
        ws.append([f"{device}: passed Core Web Vitals", f"{(ok['CWV'] == 'Passed').sum()} of {len(values)}"])
        ws.append([f"{device}: failed Core Web Vitals", int((ok["CWV"] == "Failed").sum())])
        ws.append([f"{device}: no real-user data", int((ok["CWV"] == "No data").sum())])
        perf = pd.to_numeric(ok["Performance"], errors="coerce").dropna()
        ws.append([f"{device}: average performance score", round(perf.mean()) if len(perf) else None])
    ws.append([])
    ws.append(["How to read this"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for line in (
        "Field metrics come from real Chrome users (Chrome UX Report): 75th percentile over the last 28 days.",
        "Field data 'Origin' means the URL lacks enough traffic, so the numbers are for the whole domain.",
        "Lab metrics and scores come from one Lighthouse test run by Google and can vary between runs.",
        "Core Web Vitals pass when LCP <= 2.5 s, INP <= 200 ms and CLS <= 0.1 for real users.",
    ):
        ws.append([line])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 24

    for strategy in audit["strategies"]:
        values, ratings = device_frames(audit, strategy)
        write_sheet(wb.create_sheet(DEVICE_LABEL[strategy]), values, ratings, wide=("URL", "Top fixes", "Redirects to", "Status"))

    if len(audit["strategies"]) > 1:
        values, ratings = compare_frames(audit)
        write_sheet(wb.create_sheet("Compare"), values, ratings, wide=("URL",))

    flags = collect_flags(audit)
    rows = [{"Flag": "Error", "URL": f["URL"], "Device": f["Device"], "Detail": f["Problem"]} for f in flags["errors"]]
    rows += [{"Flag": "Origin-level field data", "URL": f["URL"], "Device": f["Device"], "Detail": "Not enough traffic for this URL; domain-wide numbers shown."} for f in flags["origin"]]
    rows += [{"Flag": "No real-user data", "URL": f["URL"], "Device": f["Device"], "Detail": "Only lab data available."} for f in flags["none"]]
    rows += [{"Flag": "Redirect", "URL": u, "Device": "", "Detail": f"Redirects to {final}"} for u, final in flags["redirects"].items()]
    if rows:
        write_sheet(wb.create_sheet("Flags"), pd.DataFrame(rows), None, wide=("URL", "Detail"))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(audit: dict) -> bytes:
    frames = []
    for strategy in audit["strategies"]:
        values, _ = device_frames(audit, strategy)
        values.insert(1, "Device", DEVICE_LABEL[strategy])
        frames.append(values)
    return pd.concat(frames, ignore_index=True).to_csv(index=False).encode("utf-8-sig")


# ---------------------------------------------------------------------------
# Results: page section
# ---------------------------------------------------------------------------

def render_results(audit: dict, prefix: str) -> None:
    total = len(audit["urls"]) * len(audit["strategies"])
    done = len(audit["results"])
    st.subheader("Results", anchor=False)

    if audit["fatal"]:
        st.error(f"The audit stopped: {audit['fatal']}")
    elif audit["stopped"]:
        st.warning(f"The audit was stopped after {done} of {total} checks. Results so far are below; run it again to check the rest.")
    elif audit["finished"] and audit["duration"]:
        st.caption(f"{total} checks finished in {fmt_duration(audit['duration'])} on {audit['started'].strftime('%d %b %Y, %H:%M')} UTC.")

    st.markdown(summary_html(audit), unsafe_allow_html=True)

    flags = collect_flags(audit)
    flag_count = len(flags["errors"]) + len(flags["origin"]) + len(flags["none"]) + len(flags["redirects"])
    tab_names = [DEVICE_LABEL[s] for s in audit["strategies"]]
    if len(audit["strategies"]) > 1:
        tab_names.append("Mobile vs desktop")
    tab_names.append(f"Flags ({flag_count})")
    tabs = st.tabs(tab_names)

    for tab, strategy in zip(tabs, audit["strategies"]):
        with tab:
            values, ratings = device_frames(audit, strategy)
            show_table(values, ratings)
    if len(audit["strategies"]) > 1:
        with tabs[len(audit["strategies"])]:
            values, ratings = compare_frames(audit)
            show_table(values, ratings)
    with tabs[-1]:
        if not flag_count:
            st.write("Nothing to flag. Every URL returned data without errors or redirects.")
        if flags["errors"]:
            st.markdown("**Couldn't be checked.** Fix the cause, then run these URLs again.")
            st.dataframe(pd.DataFrame(flags["errors"]), hide_index=True, width="stretch",
                         column_config={"URL": st.column_config.LinkColumn("URL")})
        if flags["origin"]:
            st.markdown("**Field data is domain-wide.** These URLs don't get enough Chrome traffic on their own, "
                        "so the real-user numbers describe the whole domain. Use the lab metrics to judge the page itself.")
            st.dataframe(pd.DataFrame(flags["origin"]), hide_index=True, width="stretch",
                         column_config={"URL": st.column_config.LinkColumn("URL")})
        if flags["none"]:
            st.markdown("**No real-user data.** Neither the URL nor its domain has enough Chrome traffic, "
                        "so only lab results are available and the Core Web Vitals assessment can't be made.")
            st.dataframe(pd.DataFrame(flags["none"]), hide_index=True, width="stretch",
                         column_config={"URL": st.column_config.LinkColumn("URL")})
        if flags["redirects"]:
            st.markdown("**These URLs redirect.** Audit the final URL, and point internal links, canonicals and sitemaps straight to it.")
            st.dataframe(pd.DataFrame([{"URL": u, "Redirects to": f} for u, f in flags["redirects"].items()]),
                         hide_index=True, width="stretch",
                         column_config={"URL": st.column_config.LinkColumn("URL"),
                                        "Redirects to": st.column_config.LinkColumn("Redirects to")})

    if done:
        stamp = audit["started"].strftime("%Y-%m-%d")
        cache_key = (done, audit["finished"], audit["stopped"], audit["fatal"])
        if audit.get("_export_key") != cache_key:
            audit["_xlsx"], audit["_csv"] = build_excel(audit), build_csv(audit)
            audit["_export_key"] = cache_key
        c1, c2, _ = st.columns([1.3, 1.1, 2.6])
        c1.download_button(
            "Download Excel report",
            data=audit["_xlsx"],
            file_name=f"core-web-vitals-{stamp}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            width="stretch",
            on_click="ignore",
            key=f"{prefix}_xlsx",
        )
        c2.download_button(
            "Download CSV",
            data=audit["_csv"],
            file_name=f"core-web-vitals-{stamp}.csv",
            mime="text/csv",
            width="stretch",
            on_click="ignore",
            key=f"{prefix}_csv",
        )
        st.caption("The Excel report has Summary, Mobile, Desktop, comparison and Flags sheets, and opens in Excel or Google Sheets.")


# ---------------------------------------------------------------------------
# Google Sheets sync
# ---------------------------------------------------------------------------

def open_worksheet(creds: dict, sheet_ref: str, tab_name: str):
    import gspread

    client = gspread.service_account_from_dict(creds)
    ref = sheet_ref.strip()
    if "docs.google.com" in ref:
        spreadsheet = client.open_by_url(ref)
    elif re.fullmatch(r"[A-Za-z0-9_-]{30,}", ref):
        spreadsheet = client.open_by_key(ref)
    else:
        spreadsheet = client.open(ref)
    if tab_name.strip():
        return spreadsheet.worksheet(tab_name.strip())
    gid = re.search(r"[#&?]gid=(\d+)", ref)
    if gid:
        try:
            return spreadsheet.get_worksheet_by_id(int(gid.group(1)))
        except Exception:
            pass
    return spreadsheet.get_worksheet(0)


def explain_sheet_error(exc: Exception, email: str) -> str:
    import gspread

    if isinstance(exc, gspread.exceptions.SpreadsheetNotFound):
        return f"Sheet not found. Share it with {email} as Editor, and check the link."
    if isinstance(exc, gspread.exceptions.WorksheetNotFound):
        return "That tab name doesn't exist in the sheet. Check the spelling, or leave it empty to use the first tab."
    if isinstance(exc, gspread.exceptions.APIError):
        text = str(exc).lower()
        if "has not been used" in text or "disabled" in text:
            if "drive" in text:
                return "Opening a sheet by name needs the Google Drive API. Paste the sheet's link instead, or enable the Drive API in the service account's project."
            return "The Google Sheets API isn't enabled in the service account's Google Cloud project. Enable it, wait a couple of minutes, then try again."
        if "permission" in text or "403" in text:
            return f"The service account can't edit this sheet. Share it with {email} as Editor."
        return f"Google Sheets returned an error: {exc}"
    return f"Couldn't open the sheet: {exc}"


def write_to_sheet(worksheet, sheet_rows: list[str | None], audit: dict) -> int:
    from gspread.utils import rowcol_to_a1

    headers = [f"{DEVICE_LABEL[s]} {f}" for s in audit["strategies"] for f in SHEET_FIELDS] + ["Notes", "Checked (UTC)"]
    last_col = 1 + len(headers)
    if worksheet.col_count < last_col:
        worksheet.add_cols(last_col - worksheet.col_count)
    end_letter = re.sub(r"\d", "", rowcol_to_a1(1, last_col))
    stamp = audit["started"].strftime("%Y-%m-%d %H:%M")
    data = [{"range": f"B1:{end_letter}1", "values": [headers]}]
    written = 0
    for row_number, url in enumerate(sheet_rows, start=2):
        if not url or not any((url, s) in audit["results"] for s in audit["strategies"]):
            continue  # leave rows without a URL or result untouched
        values, notes = [], []
        for strategy in audit["strategies"]:
            result = audit["results"].get((url, strategy)) or {}
            device_values = result.get("values") or {}
            values += [_cell(device_values.get(f)) if device_values.get(f) is not None else "" for f in SHEET_FIELDS]
            if result.get("error"):
                notes.append(f"{DEVICE_LABEL[strategy]}: {result['error']}")
            elif not result:
                notes.append(f"{DEVICE_LABEL[strategy]}: not checked")
            if device_values.get("Redirects to"):
                notes.append(f"Redirects to {device_values['Redirects to']}")
        notes = list(dict.fromkeys(notes))
        data.append({"range": f"B{row_number}:{end_letter}{row_number}", "values": [values + ["; ".join(notes), stamp]]})
        written += 1
    worksheet.batch_update(data, value_input_option="RAW")
    return written


def read_service_account(uploaded) -> tuple[dict | None, str | None]:
    try:
        creds = json.loads(uploaded.getvalue().decode("utf-8"))
    except Exception:
        return None, "That file isn't valid JSON. Upload the .json key you downloaded for the service account."
    if not isinstance(creds, dict) or creds.get("type") != "service_account" or not creds.get("client_email") or not creds.get("private_key"):
        return None, "That file isn't a service account key. In Google Cloud Console, open the service account > Keys > Add key > JSON."
    return creds, None


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------

def thresholds_html() -> str:
    items = [("LCP", "Loading", ("≤ 2.5 s", "≤ 4 s", "> 4 s")),
             ("INP", "Responsiveness", ("≤ 200 ms", "≤ 500 ms", "> 500 ms")),
             ("CLS", "Visual stability", ("≤ 0.1", "≤ 0.25", "> 0.25"))]
    blocks = "".join(
        f"<div><div class='th-head'><b>{name}</b><span>{label}</span></div>"
        f"<div class='th-bar'><i class='g'></i><i class='n'></i><i class='p'></i></div>"
        f"<div class='th-scale'><span>{good}</span><span>{ni}</span><span>{poor}</span></div></div>"
        for name, label, (good, ni, poor) in items
    )
    note = "<div class='th-note'>Google's Core Web Vitals thresholds: good, needs improvement, poor. Measured at the 75th percentile of real visits.</div>"
    return f"<div class='thresholds' role='img' aria-label='Core Web Vitals thresholds'>{blocks}{note}</div>"


st.markdown(
    "<div class='hero'><div class='hero-copy'>"
    "<div class='hero-title' role='heading' aria-level='1'>Core Web Vitals bulk checker</div>"
    f"<p>Check up to {MAX_URLS} URLs at once on mobile and desktop. You get real-user Core Web Vitals from Chrome users, "
    "plus Lighthouse scores for performance, accessibility, best practices, SEO and agentic browsing.</p>"
    f"</div>{thresholds_html()}</div>",
    unsafe_allow_html=True,
)

# --- API key -----------------------------------------------------------------
with st.container(border=True):
    key_col, help_col = st.columns([1.2, 1], gap="large")
    with key_col:
        st.subheader("Your PageSpeed Insights API key", anchor=False)
        api_key = st.text_input(
            "API key",
            type="password",
            placeholder="Paste your key (it starts with AIza)",
            key="psi_key",
            autocomplete="off",
            label_visibility="collapsed",
        ).strip()
        if not api_key:
            st.markdown("<span class='pill idle'>No key added yet</span>", unsafe_allow_html=True)
        elif not KEY_PATTERN.match(api_key):
            st.markdown("<span class='pill warn'>This doesn't look like a Google API key. Keys start with AIza and are 39 characters long.</span>", unsafe_allow_html=True)
        else:
            st.markdown("<span class='pill ok'>Key added. Google confirms it on the first check.</span>", unsafe_allow_html=True)
    with help_col:
        st.write(
            "The checker runs on your own free Google API key, so your checks use your quota. "
            "The key is used only to call Google's PageSpeed Insights API during this visit. It isn't saved or shared, "
            "and it's gone when you close the tab."
        )
        with st.expander("How to get a free API key"):
            st.markdown(
                "1. Open Google's [PageSpeed Insights API get-started page](https://developers.google.com/speed/docs/insights/v5/get-started) "
                "and click **Get a Key**, or create one under [Credentials in Google Cloud Console](https://console.cloud.google.com/apis/credentials).\n"
                "2. Pick or create a Google Cloud project. Make sure **PageSpeed Insights API** is enabled for it.\n"
                "3. Copy the key and paste it above.\n\n"
                "**Recommended restrictions:** under API restrictions, allow only *PageSpeed Insights API*. "
                "Leave Application restrictions on *None*: a website (HTTP referrer) restriction blocks this checker."
            )

st.write("")
bulk_tab, sheet_tab = st.tabs(["Bulk check", "Google Sheets sync"])

# --- Bulk check ----------------------------------------------------------------
with bulk_tab:
    urls_col, settings_col = st.columns([1.55, 1], gap="large")
    with urls_col:
        url_text = st.text_area(
            "URLs to check",
            height=210,
            placeholder="One URL per line\nhttps://example.com/\nhttps://example.com/services/",
            key="bulk_urls",
        )
        upload = st.file_uploader("Or upload a .txt or .csv file of URLs", type=["txt", "csv"], key="bulk_upload")
        tokens = tokens_from_text(url_text) + (tokens_from_upload(upload) if upload else [])
        bulk_urls, bulk_invalid, bulk_dupes = clean_urls(tokens)
        notes = []
        if bulk_dupes:
            notes.append(f"{bulk_dupes} duplicate{'s' if bulk_dupes > 1 else ''} removed")
        if bulk_invalid:
            notes.append(f"{len(bulk_invalid)} line{'s' if len(bulk_invalid) > 1 else ''} skipped (not a URL)")
        if bulk_urls:
            st.caption(f"**{len(bulk_urls)} URL{'s' if len(bulk_urls) != 1 else ''} ready**" + (f". {_sentence(notes)}" if notes else "."))
        elif notes:
            st.caption(_sentence(notes))
        if bulk_invalid:
            with st.expander("Show skipped lines"):
                st.code("\n".join(bulk_invalid[:50]), language=None)
        if len(bulk_urls) > MAX_URLS:
            st.warning(f"Only the first {MAX_URLS} URLs will be checked in one run.")

    with settings_col:
        bulk_devices = st.radio("Devices", list(DEVICE_OPTIONS), key="bulk_devices")
        bulk_workers = st.select_slider(
            "Checks at a time",
            options=[1, 2, 3, 4, 5, 6],
            value=4,
            key="bulk_workers",
            help="More is faster. Lower it if Google starts rate-limiting your key.",
        )
        strategies = DEVICE_OPTIONS[bulk_devices]
        checks = min(len(bulk_urls), MAX_URLS) * len(strategies)
        if checks:
            st.caption(f"{checks} checks, about {estimate_minutes(checks, bulk_workers)} min.")
        run_bulk = st.button("Run audit", type="primary", width="stretch", key="bulk_run")

    if run_bulk:
        if not api_key:
            st.error("Add your PageSpeed Insights API key above to run the audit.")
        elif not KEY_PATTERN.match(api_key):
            st.error("That API key doesn't look right. Copy the whole key from Google Cloud Console; it starts with AIza and is 39 characters long.")
        elif not bulk_urls:
            st.error("Add at least one URL, one per line, or upload a .txt or .csv file.")
        else:
            audit = new_audit(bulk_urls[:MAX_URLS], strategies)
            st.session_state["bulk_audit"] = audit
            execute_audit(audit, api_key, bulk_workers, "bulk")

    if st.session_state.get("bulk_audit"):
        render_results(st.session_state["bulk_audit"], "bulk")

# --- Google Sheets sync ----------------------------------------------------------
with sheet_tab:
    st.write(
        "Reads URLs from column A of your Google Sheet (from row 2), checks them, and writes the results into "
        "the columns next to them. Columns B onwards are overwritten for rows that get a result."
    )
    with st.expander("One-time setup: connect your Google Sheet"):
        st.markdown(
            "1. In [Google Cloud Console](https://console.cloud.google.com/), open a project and enable the **Google Sheets API**.\n"
            "2. Go to **IAM & Admin > Service accounts**, create a service account, then open it and choose "
            "**Keys > Add key > Create new key > JSON**. A .json file downloads.\n"
            "3. Upload that file below. The app shows the service account's email address.\n"
            "4. In your Google Sheet, click **Share**, paste that email and give it **Editor** access.\n"
            "5. Put your URLs in column A from row 2, paste the sheet link below and run the sync.\n\n"
            "The key file is used only during this visit and isn't saved. If your organisation blocks service account keys, "
            "use **Bulk check** and open the Excel report in Google Sheets instead."
        )

    sheet_col, creds_col = st.columns(2, gap="large")
    with sheet_col:
        sheet_ref = st.text_input("Google Sheet link", placeholder="https://docs.google.com/spreadsheets/d/…", key="sheet_ref")
        sheet_tab_name = st.text_input("Tab name (optional)", placeholder="Leave empty to use the tab in the link, or the first tab", key="sheet_tab")
    with creds_col:
        sa_file = st.file_uploader("Service account key (.json)", type=["json"], key="sheet_creds")
        sa_creds, sa_error = (read_service_account(sa_file) if sa_file else (None, None))
        if sa_error:
            st.error(sa_error)
        elif sa_creds:
            st.markdown("Share your sheet with this email as **Editor**:")
            st.code(sa_creds["client_email"], language=None)

    s1, s2, s3 = st.columns([1.2, 1, 1], vertical_alignment="bottom")
    sheet_devices = s1.radio("Devices", list(DEVICE_OPTIONS), key="sheet_devices", horizontal=True)
    sheet_workers = s2.select_slider("Checks at a time", options=[1, 2, 3, 4, 5, 6], value=4, key="sheet_workers")
    run_sheet = s3.button("Run sync", type="primary", width="stretch", key="sheet_run")

    if run_sheet:
        if not api_key:
            st.error("Add your PageSpeed Insights API key above to run the sync.")
        elif not KEY_PATTERN.match(api_key):
            st.error("That API key doesn't look right. Copy the whole key from Google Cloud Console; it starts with AIza and is 39 characters long.")
        elif not sheet_ref.strip():
            st.error("Paste your Google Sheet link.")
        elif not sa_creds:
            st.error("Upload your service account key (.json) so the app can open the sheet.")
        else:
            try:
                with st.spinner("Opening your sheet…"):
                    worksheet = open_worksheet(sa_creds, sheet_ref, sheet_tab_name)
                    column_a = worksheet.col_values(1)[1:]
            except Exception as exc:
                st.error(explain_sheet_error(exc, sa_creds["client_email"]))
            else:
                sheet_rows = [normalise_url(cell) for cell in column_a]
                unique_urls = list(dict.fromkeys(u for u in sheet_rows if u))
                if not unique_urls:
                    st.error(f"No URLs found in column A of “{worksheet.title}” (row 1 is treated as a header).")
                else:
                    if len(unique_urls) > MAX_URLS:
                        st.warning(f"The sheet has {len(unique_urls)} URLs. The first {MAX_URLS} will be checked in this run.")
                    audit = new_audit(unique_urls[:MAX_URLS], DEVICE_OPTIONS[sheet_devices])
                    audit["sheet"] = {"creds": sa_creds, "ref": sheet_ref, "tab": sheet_tab_name, "rows": sheet_rows, "title": worksheet.title}
                    st.session_state["sheet_audit"] = audit
                    st.caption(f"Found {len(unique_urls)} URLs in “{worksheet.title}”.")
                    execute_audit(audit, api_key, sheet_workers, "sheet")
                    try:
                        with st.spinner("Writing results to your sheet…"):
                            written = write_to_sheet(worksheet, sheet_rows, audit)
                        audit["sheet"]["written"] = written
                    except Exception as exc:
                        audit["sheet"]["write_error"] = explain_sheet_error(exc, sa_creds["client_email"])

    sheet_audit = st.session_state.get("sheet_audit")
    if sheet_audit:
        info = sheet_audit["sheet"]
        if info.get("write_error"):
            st.error(f"Results couldn't be written to the sheet. {info['write_error']} You can still download them below.")
        elif info.get("written") is not None:
            st.success(f"Updated {info['written']} rows in “{info['title']}”.")
        elif sheet_audit["results"]:
            if st.button(f"Write the {len(sheet_audit['results'])} finished checks to the sheet", key="sheet_write_partial"):
                try:
                    worksheet = open_worksheet(info["creds"], info["ref"], info["tab"])
                    info["written"] = write_to_sheet(worksheet, info["rows"], sheet_audit)
                    st.rerun()
                except Exception as exc:
                    info["write_error"] = explain_sheet_error(exc, info["creds"]["client_email"])
                    st.rerun()
        render_results(sheet_audit, "sheet")

st.markdown(
    "<div class='footnote'>Data comes from Google's PageSpeed Insights API. Field metrics are from the Chrome UX Report: "
    "real Chrome users over the last 28 days, at the 75th percentile. Lab metrics and scores come from a single Lighthouse "
    "run, so they can shift a few points between runs. Agentic browsing is an experimental Lighthouse category.</div>",
    unsafe_allow_html=True,
)
